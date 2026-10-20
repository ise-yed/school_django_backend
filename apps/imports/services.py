import io
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Font
from django.db import transaction
from rapidfuzz import fuzz, process
from rest_framework.exceptions import ValidationError

from apps.accounts.models import User
from apps.core.normalizers import (
    normalize_and_validate_national_code,
    normalize_digits,
    normalize_persian_text,
    normalize_phone,
)
from apps.imports.models import ImportLog, ImportMapping
from apps.schools.models import Classroom
from apps.students.models import Enrollment, ParentStudentLink, Student
from apps.students.services import build_normalized_search_name, split_persian_full_name

# دیکشنری مترادف‌های ستون‌های اکسل مدارس
COLUMN_SYNONYMS: dict[str, list[str]] = {
    "full_name": [
        "اسم",
        "نام و نام خانوادگی",
        "نام دانش اموز",
        "نام کامل",
        "مشخصات دانش اموز",
    ],
    "first_name": [
        "نام",
        "اسم کوچک",
    ],
    "last_name": [
        "نام خانوادگی",
        "فامیل",
        "شهرت",
    ],
    "father_name": [
        "نام پدر",
        "اسم پدر",
        "پدر",
        "فرزند",
    ],
    "national_code": [
        "کد ملی",
        "کدملی",
        "شماره ملی",
        "شناسه ملی",
        "کد ملی دانش اموز",
    ],
    "parent_phone": [
        "موبایل ولی",
        "شماره همراه",
        "تلفن همراه",
        "موبایل",
        "شماره موبایل",
        "شماره ولی",
        "تلفن ولی",
        "همراه ولی",
    ],
    "student_code": [
        "کد دانش اموزی",
        "شماره دانش اموزی",
        "کد سناد",
    ],
}


def _cell_to_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def load_workbook_safely(file_obj):
    filename = getattr(file_obj, "name", "") or ""
    ext = Path(filename).suffix.lower()
    if ext and ext not in (".xlsx", ".xlsm"):
        raise ValidationError(
            "فرمت فایل پشتیبانی نمی‌شود. لطفاً فایل اکسل استاندارد با فرمت .xlsx آپلود کنید."
        )
    try:
        file_obj.seek(0)
        return openpyxl.load_workbook(file_obj, data_only=True)
    except Exception:
        raise ValidationError(
            "فایل اکسل قابل خواندن نیست یا آسیب دیده است. لطفاً یک فایل معتبر .xlsx انتخاب کنید."
        )


def unmerge_and_extract_rows(ws) -> list[list[str]]:
    """
    باز کردن سلول‌های ادغام‌شده (Merged Cells) و استخراج تمام ردیف‌ها به‌صورت رشته.
    """
    for merged_range in list(ws.merged_cells.ranges):
        top_left_value = ws.cell(merged_range.min_row, merged_range.min_col).value
        ws.unmerge_cells(str(merged_range))
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                ws.cell(row, col).value = top_left_value

    rows = []
    for row in ws.iter_rows(values_only=True):
        str_row = [_cell_to_str(cell) for cell in row]
        rows.append(str_row)
    return rows


def guess_header_row_and_columns(rows: list[list[str]]) -> tuple[int, list[str]]:
    """
    تشخیص خودکار ردیف سرستون در ۱۵ ردیف اول فایل اکسل (نادیده گرفتن عنوان بالای جدول یا ردیف خالی).
    شماره ردیف به‌صورت 1-based برگردانده می‌شود.
    """
    all_synonyms = [
        normalize_persian_text(syn)
        for syns in COLUMN_SYNONYMS.values()
        for syn in syns
    ]

    best_row_idx = 0
    best_score = -1

    for idx, row in enumerate(rows[:15]):
        non_empty = [normalize_persian_text(c) for c in row if normalize_persian_text(c)]
        if len(non_empty) < 2:
            continue
        score = 0
        for cell in non_empty:
            if cell in all_synonyms:
                score += 3
            else:
                match = process.extractOne(cell, all_synonyms, scorer=fuzz.ratio)
                if match and match[1] >= 80:
                    score += 2
        if score > best_score:
            best_score = score
            best_row_idx = idx

    if not rows:
        return 1, []

    raw_headers = rows[best_row_idx]
    headers = [
        normalize_persian_text(h) if normalize_persian_text(h) else f"ستون_{i + 1}"
        for i, h in enumerate(raw_headers)
    ]
    return best_row_idx + 1, headers


def suggest_column_mapping(school, headers: list[str]) -> dict:
    """
    پیشنهاد نگاشت ستون‌ها با اولویت:
    ۱. الگوی ذخیره‌شده مدرسه در ImportMapping
    ۲. تطبیق دقیق با دیکشنری مترادف‌ها پس از نرمال‌سازی فارسی
    ۳. تطبیق فازی با rapidfuzz (آستانه >= 80)
    """
    saved_mapping = ImportMapping.objects.for_school(school).first()
    mapping = {
        "first_name": None,
        "last_name": None,
        "full_name": None,
        "father_name": None,
        "national_code": None,
        "parent_phone": None,
        "student_code": None,
        "split_full_name": False,
    }

    normalized_headers = {normalize_persian_text(h): h for h in headers}

    if saved_mapping and isinstance(saved_mapping.mapping_config, dict):
        for field_key in mapping.keys():
            if field_key == "split_full_name":
                continue
            saved_col = saved_mapping.mapping_config.get(field_key)
            if saved_col and normalize_persian_text(saved_col) in normalized_headers:
                mapping[field_key] = normalized_headers[normalize_persian_text(saved_col)]

    used_headers = {v for k, v in mapping.items() if k != "split_full_name" and v}

    # مرحله ۱: تطبیق دقیق
    for field_key, synonyms in COLUMN_SYNONYMS.items():
        if mapping.get(field_key):
            continue
        norm_syns = [normalize_persian_text(s) for s in synonyms]
        for norm_h, orig_h in normalized_headers.items():
            if orig_h in used_headers:
                continue
            if norm_h in norm_syns:
                mapping[field_key] = orig_h
                used_headers.add(orig_h)
                break

    # مرحله ۲: تطبیق فازی با rapidfuzz
    for field_key, synonyms in COLUMN_SYNONYMS.items():
        if mapping.get(field_key):
            continue
        norm_syns = [normalize_persian_text(s) for s in synonyms]
        best_header = None
        best_score = 0
        for norm_h, orig_h in normalized_headers.items():
            if orig_h in used_headers:
                continue
            match = process.extractOne(norm_h, norm_syns, scorer=fuzz.ratio)
            if match and match[1] >= 80 and match[1] > best_score:
                best_score = match[1]
                best_header = orig_h
        if best_header:
            mapping[field_key] = best_header
            used_headers.add(best_header)

    # اگر ستون ترکیبی «اسم» پیدا شد ولی «نام» و «نام خانوادگی» جدا وجود نداشت، گزینه تقسیم فعال شود
    if mapping["full_name"] and not (mapping["first_name"] and mapping["last_name"]):
        mapping["split_full_name"] = True

    return mapping


def guess_classroom_from_text(school, filename: str, sheet_name: str) -> dict | None:
    """حدس کلاس مقصد از روی نام فایل یا نام شیت در سال تحصیلی جاری مدرسه."""
    classrooms = list(
        Classroom.objects.for_school(school)
        .filter(academic_year__is_current=True)
        .select_related("grade", "academic_year")
    )
    if not classrooms:
        classrooms = list(
            Classroom.objects.for_school(school).select_related("grade", "academic_year")
        )
    if not classrooms:
        return None

    stem = Path(filename or "").stem.replace("_", " ").replace("-", " ")
    target_text = normalize_persian_text(f"{stem} {sheet_name or ''}")

    best_classroom = None
    best_score = 0
    for c in classrooms:
        c_name = normalize_persian_text(c.name)
        g_name = normalize_persian_text(c.grade.title)
        full_label = normalize_persian_text(f"{g_name} {c_name}")

        if c_name and c_name in target_text:
            score = 95
        else:
            score = max(
                fuzz.partial_ratio(c_name, target_text),
                fuzz.partial_ratio(full_label, target_text),
            )
        if score > best_score and score >= 70:
            best_score = score
            best_classroom = c

    if not best_classroom:
        return None
    return {
        "id": best_classroom.id,
        "name": best_classroom.name,
        "grade_title": best_classroom.grade.title,
    }


def analyze_excel_file(school, user, file_obj, sheet_name: str | None = None) -> dict:
    wb = load_workbook_safely(file_obj)
    sheets = wb.sheetnames
    selected_sheet = sheet_name if sheet_name in sheets else sheets[0]
    ws = wb[selected_sheet]

    rows = unmerge_and_extract_rows(ws)
    header_row_1based, columns = guess_header_row_and_columns(rows)
    suggested_mapping = suggest_column_mapping(school, columns)

    data_rows = rows[header_row_1based:]
    sample_rows = []
    for r in data_rows:
        if not any(c.strip() for c in r):
            continue
        row_dict = {
            columns[i]: (r[i] if i < len(r) else "")
            for i in range(len(columns))
        }
        sample_rows.append(row_dict)
        if len(sample_rows) >= 5:
            break

    file_obj.seek(0)
    import_log = ImportLog.objects.create(
        school=school,
        uploaded_by=user,
        file=file_obj,
        original_filename=getattr(file_obj, "name", ""),
        sheet_name=selected_sheet,
        header_row=header_row_1based,
        mapping_used=suggested_mapping,
        status=ImportLog.Status.ANALYZED,
    )

    suggested_classroom = guess_classroom_from_text(
        school, import_log.original_filename, selected_sheet
    )

    return {
        "import_log_id": import_log.id,
        "sheets": sheets,
        "selected_sheet": selected_sheet,
        "guessed_header_row": header_row_1based,
        "columns": columns,
        "suggested_mapping": suggested_mapping,
        "sample_rows": sample_rows,
        "suggested_classroom": suggested_classroom,
    }


def parse_and_validate_rows(
    school,
    classroom: Classroom,
    rows: list[list[str]],
    header_row_1based: int,
    mapping: dict,
    row_overrides: dict | None = None,
) -> dict:
    """
    نرمال‌سازی و اعتبارسنجی ردیف‌های اکسل و مقایسه‌ی آن‌ها با دیتابیس مدرسه و کلاس مقصد.
    """
    row_overrides = row_overrides or {}
    if not rows or header_row_1based > len(rows):
        raise ValidationError("ردیف سرستون در فایل یافت نشد.")

    raw_headers = rows[header_row_1based - 1]
    headers = [
        normalize_persian_text(h) if normalize_persian_text(h) else f"ستون_{i + 1}"
        for i, h in enumerate(raw_headers)
    ]
    col_index = {h: idx for idx, h in enumerate(headers)}
    # پشتیبانی از نام ستون بدون نرمال‌سازی در صورتی که مستقیماً ارسال شده باشد
    for idx, raw_h in enumerate(raw_headers):
        if raw_h and raw_h not in col_index:
            col_index[raw_h] = idx

    def get_val(row_values: list[str], field_name: str) -> str:
        col_name = mapping.get(field_name)
        if not col_name:
            return ""
        idx = col_index.get(col_name)
        if idx is None:
            idx = col_index.get(normalize_persian_text(col_name))
        if idx is None or idx >= len(row_values):
            return ""
        return row_values[idx]

    valid_rows = []
    errors = []
    seen_national_codes: dict[str, int] = {}

    for offset, r in enumerate(rows[header_row_1based:]):
        excel_row_num = header_row_1based + 1 + offset
        if not any(c.strip() for c in r):
            continue  # نادیده گرفتن ردیف‌های کاملاً خالی

        raw_nc = get_val(r, "national_code")
        raw_first = get_val(r, "first_name")
        raw_last = get_val(r, "last_name")
        raw_full = get_val(r, "full_name")
        raw_father = get_val(r, "father_name")
        raw_phone = get_val(r, "parent_phone")
        raw_student_code = get_val(r, "student_code")

        if mapping.get("split_full_name") or (raw_full and not (raw_first or raw_last)):
            first_name, last_name = split_persian_full_name(raw_full)
        else:
            first_name = normalize_persian_text(raw_first)
            last_name = normalize_persian_text(raw_last)

        # اعمال اصلاحات دستی کاربر در پیش‌نمایش (در صورت وجود)
        override = row_overrides.get(str(excel_row_num)) or row_overrides.get(excel_row_num)
        if isinstance(override, dict):
            if "first_name" in override:
                first_name = normalize_persian_text(override["first_name"])
            if "last_name" in override:
                last_name = normalize_persian_text(override["last_name"])

        father_name = normalize_persian_text(raw_father)
        student_code = normalize_digits(raw_student_code) or None

        row_errors = []
        national_code = ""
        if not raw_nc:
            row_errors.append("کد ملی خالی است.")
        else:
            try:
                national_code = normalize_and_validate_national_code(raw_nc)
            except ValueError as e:
                row_errors.append(str(e))

        if not first_name:
            row_errors.append("نام دانش‌آموز الزامی است.")
        if not last_name:
            row_errors.append("نام خانوادگی دانش‌آموز الزامی است.")

        if national_code:
            if national_code in seen_national_codes:
                prev_row = seen_national_codes[national_code]
                row_errors.append(
                    f"کد ملی تکراری در داخل فایل (قبلاً در ردیف {prev_row} آمده است)."
                )
            else:
                seen_national_codes[national_code] = excel_row_num

        parent_phone = ""
        if raw_phone:
            try:
                parent_phone = normalize_phone(raw_phone)
            except ValueError as e:
                row_errors.append(f"خطای موبایل ولی: {e}")

        if row_errors:
            errors.append(
                {
                    "row_number": excel_row_num,
                    "raw_data": {
                        "full_name": raw_full or f"{raw_first} {raw_last}".strip(),
                        "father_name": raw_father,
                        "national_code": raw_nc,
                        "parent_phone": raw_phone,
                    },
                    "error_reason": " | ".join(row_errors),
                }
            )
            continue

        valid_rows.append(
            {
                "row_number": excel_row_num,
                "first_name": first_name,
                "last_name": last_name,
                "father_name": father_name,
                "national_code": national_code,
                "parent_phone": parent_phone,
                "student_code": student_code,
            }
        )

    # بررسی وضعیت دانش‌آموزان در دیتابیس مدرسه
    file_codes = [ item["national_code"] for item in valid_rows ]
    existing_students_map = {
        s.national_code: s
        for s in Student.objects.for_school(school).filter(national_code__in=file_codes)
    }

    existing_enrollments_map = {
        e.student.national_code: e
        for e in Enrollment.objects.for_school(school)
        .filter(
            academic_year=classroom.academic_year,
            student__national_code__in=file_codes,
        )
        .select_related("classroom", "student")
    }

    new_students = []
    existing_students = []

    for item in valid_rows:
        nc = item["national_code"]
        if nc in existing_students_map:
            current_enrollment = existing_enrollments_map.get(nc)
            requires_transfer = bool(
                current_enrollment and current_enrollment.classroom_id != classroom.id
            )
            existing_students.append(
                {
                    **item,
                    "student_id": existing_students_map[nc].id,
                    "previous_classroom_id": (
                        current_enrollment.classroom_id if current_enrollment else None
                    ),
                    "previous_classroom_name": (
                        current_enrollment.classroom.name if current_enrollment else None
                    ),
                    "requires_classroom_transfer": requires_transfer,
                }
            )
        else:
            new_students.append(item)

    # دانش‌آموزانی که هم‌اکنون در این کلاس ثبت‌نام دارند اما در فایل اکسل نیستند (فقط گزارش؛ بدون حذف خودکار)
    missing_enrollments = (
        Enrollment.objects.for_school(school)
        .filter(classroom=classroom, is_active=True)
        .exclude(student__national_code__in=file_codes)
        .select_related("student")
    )
    missing_from_file_students = [
        {
            "student_id": e.student.id,
            "first_name": e.student.first_name,
            "last_name": e.student.last_name,
            "father_name": e.student.father_name,
            "national_code": e.student.national_code,
        }
        for e in missing_enrollments
    ]

    return {
        "total_rows": len(valid_rows) + len(errors),
        "valid_count": len(valid_rows),
        "error_count": len(errors),
        "errors": errors,
        "valid_rows": valid_rows,
        "new_students": new_students,
        "existing_students": existing_students,
        "missing_from_file_students": missing_from_file_students,
    }


def preview_import(
    import_log: ImportLog,
    classroom: Classroom,
    mapping: dict,
    sheet_name: str | None = None,
    header_row: int | None = None,
    row_overrides: dict | None = None,
) -> dict:
    wb = load_workbook_safely(import_log.file)
    selected_sheet = sheet_name or import_log.sheet_name or wb.sheetnames[0]
    ws = wb[selected_sheet]
    rows = unmerge_and_extract_rows(ws)
    used_header_row = header_row or import_log.header_row or 1

    result = parse_and_validate_rows(
        school=import_log.school,
        classroom=classroom,
        rows=rows,
        header_row_1based=used_header_row,
        mapping=mapping,
        row_overrides=row_overrides,
    )

    import_log.target_classroom = classroom
    import_log.sheet_name = selected_sheet
    import_log.header_row = used_header_row
    import_log.mapping_used = mapping
    import_log.status = ImportLog.Status.PREVIEWED
    import_log.total_rows = result["total_rows"]
    import_log.error_count = result["error_count"]
    import_log.errors_detail = result["errors"]
    import_log.save()

    return {
        "import_log_id": import_log.id,
        "classroom_id": classroom.id,
        "classroom_name": classroom.name,
        "total_rows": result["total_rows"],
        "valid_count": result["valid_count"],
        "error_count": result["error_count"],
        "errors": result["errors"],
        "new_students": result["new_students"],
        "existing_students": result["existing_students"],
        "missing_from_file_students": result["missing_from_file_students"],
    }


def commit_import(
    import_log: ImportLog,
    classroom: Classroom,
    mapping: dict,
    allow_classroom_transfer: bool = False,
    save_mapping: bool = False,
    mapping_name: str = "پیش‌فرض",
    sheet_name: str | None = None,
    header_row: int | None = None,
    row_overrides: dict | None = None,
) -> dict:
    school = import_log.school
    wb = load_workbook_safely(import_log.file)
    selected_sheet = sheet_name or import_log.sheet_name or wb.sheetnames[0]
    ws = wb[selected_sheet]
    rows = unmerge_and_extract_rows(ws)
    used_header_row = header_row or import_log.header_row or 1

    parsed = parse_and_validate_rows(
        school=school,
        classroom=classroom,
        rows=rows,
        header_row_1based=used_header_row,
        mapping=mapping,
        row_overrides=row_overrides,
    )

    valid_rows = parsed["valid_rows"]
    file_codes = [r["national_code"] for r in valid_rows]

    with transaction.atomic():
        existing_map = {
            s.national_code: s
            for s in Student.objects.for_school(school)
            .select_for_update()
            .filter(national_code__in=file_codes)
        }

        to_create_students = []
        to_update_students = []

        for r in valid_rows:
            nc = r["national_code"]
            norm_name = build_normalized_search_name(
                r["first_name"], r["last_name"], r["father_name"]
            )
            if nc in existing_map:
                st = existing_map[nc]
                st.first_name = r["first_name"]
                st.last_name = r["last_name"]
                st.father_name = r["father_name"]
                st.normalized_search_name = norm_name
                if r["student_code"]:
                    st.student_code = r["student_code"]
                st.is_active = True
                to_update_students.append(st)
            else:
                to_create_students.append(
                    Student(
                        school=school,
                        first_name=r["first_name"],
                        last_name=r["last_name"],
                        father_name=r["father_name"],
                        national_code=nc,
                        student_code=r["student_code"],
                        normalized_search_name=norm_name,
                        is_active=True,
                    )
                )

        if to_create_students:
            Student.objects.bulk_create(to_create_students)

        if to_update_students:
            Student.objects.bulk_update(
                to_update_students,
                [
                    "first_name",
                    "last_name",
                    "father_name",
                    "normalized_search_name",
                    "student_code",
                    "is_active",
                ],
            )

        # واکشی مجدد همه‌ی دانش‌آموزان فایل برای ثبت‌نام کلاس و اتصال ولی
        all_students_map = {
            s.national_code: s
            for s in Student.objects.for_school(school).filter(national_code__in=file_codes)
        }

        existing_enrollments_map = {
            e.student_id: e
            for e in Enrollment.objects.for_school(school)
            .select_for_update()
            .filter(
                academic_year=classroom.academic_year,
                student__in=all_students_map.values(),
            )
        }

        enrollments_to_create = []
        enrollments_to_update = []
        skipped_transfers = 0

        for r in valid_rows:
            st = all_students_map[r["national_code"]]
            enr = existing_enrollments_map.get(st.id)
            if enr is None:
                enrollments_to_create.append(
                    Enrollment(
                        school=school,
                        student=st,
                        academic_year=classroom.academic_year,
                        classroom=classroom,
                        is_active=True,
                    )
                )
            elif enr.classroom_id != classroom.id:
                if allow_classroom_transfer:
                    enr.classroom = classroom
                    enr.is_active = True
                    enrollments_to_update.append(enr)
                else:
                    skipped_transfers += 1

        if enrollments_to_create:
            Enrollment.objects.bulk_create(enrollments_to_create)
        if enrollments_to_update:
            Enrollment.objects.bulk_update(enrollments_to_update, ["classroom", "is_active"])

        # اگر ستون اختیاری «موبایل ولی» در اکسل پر شده بود، حساب ولی و اتصال تاییدشده ساخته شود
        for r in valid_rows:
            phone = r.get("parent_phone")
            if phone:
                st = all_students_map[r["national_code"]]
                parent_user, _ = User.objects.get_or_create(
                    school=school,
                    phone_number=phone,
                    defaults={"role": User.Role.PARENT, "is_active": True},
                )
                link, created = ParentStudentLink.objects.get_or_create(
                    school=school,
                    parent_user=parent_user,
                    student=st,
                    defaults={"is_verified": True},
                )
                if not created and not link.is_verified:
                    link.is_verified = True
                    link.save(update_fields=["is_verified"])

        if save_mapping:
            ImportMapping.objects.update_or_create(
                school=school,
                name=mapping_name or "پیش‌فرض",
                defaults={"mapping_config": mapping},
            )

        import_log.target_classroom = classroom
        import_log.mapping_used = mapping
        import_log.status = ImportLog.Status.COMMITTED
        import_log.total_rows = parsed["total_rows"]
        import_log.created_count = len(to_create_students)
        import_log.updated_count = len(to_update_students)
        import_log.error_count = parsed["error_count"]
        import_log.errors_detail = parsed["errors"]
        import_log.save()

    return {
        "import_log_id": import_log.id,
        "status": import_log.status,
        "total_rows": import_log.total_rows,
        "created_count": import_log.created_count,
        "updated_count": import_log.updated_count,
        "error_count": import_log.error_count,
        "skipped_transfers_count": skipped_transfers,
        "errors": parsed["errors"],
    }


def generate_sample_template_workbook() -> bytes:
    """
    ساخت فایل اکسل نمونه با ستون‌های «اسم»، «نام پدر»، «کد ملی»، «موبایل ولی»
    که ستون کد ملی و موبایل ولی روی فرمت متن (@) تنظیم شده‌اند.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "لیست دانش‌آموزان"
    ws. sheet_view.rightToLeft = True

    headers = ["اسم", "نام پدر", "کد ملی", "موبایل ولی"]
    ws.append(headers)

    sample_row = ["سید محمد حسینی", "علی", "0012345679", "09121112233"]
    ws.append(sample_row)

    for col_idx in range(1, 5):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # تنظیم فرمت ستون کد ملی (C) و موبایل ولی (D) روی Text (@) برای حفظ صفر ابتدایی
    for row_idx in range(2, 100):
        ws.cell(row=row_idx, column=3).number_format = "@"
        ws.cell(row=row_idx, column=4).number_format = "@"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generate_errors_workbook(import_log: ImportLog) -> bytes:
    """تولید فایل اکسل قابل دانلود از ردیف‌های دارای خطا در ImportLog."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "خطاهای ایمپورت"
    ws.sheet_view.rightToLeft = True

    headers = ["شماره ردیف در فایل", "اسم", "نام پدر", "کد ملی", "موبایل ولی", "علت خطا"]
    ws.append(headers)

    for err in import_log.errors_detail:
        raw = err.get("raw_data", {})
        ws.append(
            [
                err.get("row_number"),
                raw.get("full_name", ""),
                raw.get("father_name", ""),
                raw.get("national_code", ""),
                raw.get("parent_phone", ""),
                err.get("error_reason", ""),
            ]
        )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()